"""Document loaders for the supported input formats."""

import re
import hashlib
from dataclasses import replace
from abc import ABC, abstractmethod
from collections import Counter
from pathlib import Path
from zipfile import BadZipFile

from smart_audiobook.chapter_detection import detect_chapters
from smart_audiobook.models import Document, DocumentBlock, DocumentFormat
from smart_audiobook.text_normalizer import normalize_text


class DocumentLoadError(ValueError):
    """Raised when a source document cannot be loaded safely."""


class DocumentLoader(ABC):
    """Shared interface implemented by every source-format adapter."""

    @abstractmethod
    def load(self, path: Path) -> Document:
        """Extract and normalize one source document."""


class TxtDocumentLoader(DocumentLoader):
    def load(self, path: Path) -> Document:
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError as error:
            raise DocumentLoadError(
                "No se pudo leer el TXT. Guárdalo con codificación UTF-8."
            ) from error
        except OSError as error:
            raise DocumentLoadError(f"No se pudo leer el TXT: {error}") from error
        return _build_document(path, "txt", text)


class EpubDocumentLoader(DocumentLoader):
    def load(self, path: Path) -> Document:
        from smart_audiobook.epub_importer import load_epub
        try:
            return load_epub(path)
        except Exception as error:
            raise DocumentLoadError(f"No se pudo importar el EPUB: {error}") from error


class PdfDocumentLoader(DocumentLoader):
    def load(self, path: Path) -> Document:
        try:
            from pypdf import PdfReader

            reader = PdfReader(path)
            if reader.is_encrypted:
                try:
                    unlocked = reader.decrypt("")
                except Exception as error:
                    raise DocumentLoadError(
                        "El PDF está protegido y no se puede leer sin contraseña."
                    ) from error
                if not unlocked:
                    raise DocumentLoadError(
                        "El PDF está protegido y no se puede leer sin contraseña."
                    )
            pages = [page.extract_text() or "" for page in reader.pages]
            metadata_title = (
                (getattr(reader.metadata, "title", None) or "").strip()
                if reader.metadata
                else ""
            )
        except DocumentLoadError:
            raise
        except Exception as error:
            raise DocumentLoadError(f"No se pudo leer el PDF: {error}") from error

        if not any(page.strip() for page in pages):
            raise DocumentLoadError(
                "El PDF no contiene texto extraíble. Puede ser un documento "
                "escaneado y necesitar OCR, que todavía no forma parte de V0.4."
            )
        return _build_document(
            path,
            "pdf",
            _clean_pdf_pages(pages),
            title=metadata_title,
            author=(getattr(reader.metadata, "author", None) or None),
            source_pages=pages,
        )


class DocxDocumentLoader(DocumentLoader):
    def load(self, path: Path) -> Document:
        try:
            from docx import Document as WordDocument
            from docx.opc.exceptions import PackageNotFoundError

            word_document = WordDocument(path)
            blocks: list[DocumentBlock] = []
            title = (word_document.core_properties.title or "").strip()
            for paragraph in word_document.paragraphs:
                text = normalize_text(paragraph.text)
                if not text:
                    continue
                style_name = paragraph.style.name if paragraph.style else ""
                if style_name.casefold() == "title":
                    if not title:
                        title = text
                    continue
                blocks.append(
                    DocumentBlock(
                        text=text,
                        heading_level=_heading_level(style_name),
                        emphasis=tuple(sorted({kind for run in paragraph.runs for kind, enabled in
                            (("italic", run.italic), ("bold", run.bold)) if enabled})),
                        source_reference=f"docx:paragraph-{len(blocks)+1}",
                    )
                )
        except (PackageNotFoundError, BadZipFile, OSError, ValueError) as error:
            raise DocumentLoadError(f"No se pudo leer el DOCX: {error}") from error
        except Exception as error:
            raise DocumentLoadError(f"No se pudo procesar el DOCX: {error}") from error

        full_text = normalize_text("\n\n".join(block.text for block in blocks))
        if not full_text:
            raise DocumentLoadError("El documento DOCX está vacío.")
        document = Document(
            title=title or path.stem,
            source_path=path.resolve(),
            format="docx",
            full_text=full_text,
            blocks=tuple(blocks),
            author=word_document.core_properties.author or None,
        )
        return _enrich_chapters(detect_chapters(document))


_LOADERS: dict[str, DocumentLoader] = {
    ".txt": TxtDocumentLoader(),
    ".pdf": PdfDocumentLoader(),
    ".docx": DocxDocumentLoader(),
    ".epub": EpubDocumentLoader(),
}


def load_document(path: Path) -> Document:
    """Load a document by extension through the common adapter interface."""
    if not path.is_file():
        raise DocumentLoadError(f"No se encontró el archivo de entrada: {path}")
    loader = _LOADERS.get(path.suffix.casefold())
    if loader is None:
        supported = ", ".join(sorted(_LOADERS))
        raise DocumentLoadError(
            f"Formato no soportado: {path.suffix or '(sin extensión)'}. "
            f"Formatos disponibles: {supported}."
        )
    return loader.load(path)


def _build_document(
    path: Path,
    document_format: DocumentFormat,
    text: str,
    title: str = "",
    author: str | None = None,
    source_pages: list[str] | None = None,
) -> Document:
    normalized = normalize_text(text)
    if not normalized:
        raise DocumentLoadError(f"El documento {document_format.upper()} está vacío.")
    blocks = tuple(
        DocumentBlock(text=line.strip())
        for line in normalized.splitlines()
        if line.strip()
    )
    if source_pages is not None:
        blocks = tuple(DocumentBlock(text=normalize_text(line), source_reference=f"pdf:page-{number}")
            for number, page in enumerate(_pdf_page_texts(source_pages), start=1)
            for line in page.splitlines() if normalize_text(line))
    document = detect_chapters(
        Document(
            title=normalize_text(title) or path.stem,
            source_path=path.resolve(),
            format=document_format,
            full_text=normalized,
            blocks=blocks,
            author=author,
        )
    )
    return _enrich_chapters(document, source_pages)


def _enrich_chapters(document: Document, pages: list[str] | None = None) -> Document:
    """Attach stable identities and conservative source references to old adapters."""
    chapters = []
    cursor = 0
    for chapter in document.chapters:
        start = document.full_text.find(chapter.text, cursor)
        start = max(cursor, start)
        end = start + len(chapter.text)
        reference = f"txt:characters-{start}-{end}"
        if pages is not None:
            hits = sorted({int(b.source_reference.rsplit("-", 1)[-1]) for b in chapter.blocks
                           if b.source_reference and b.source_reference.startswith("pdf:page-")})
            reference = f"pdf:pages-{min(hits)}-{max(hits)}" if hits else "pdf:pages-unknown"
        blocks = chapter.blocks or tuple(DocumentBlock(line) for line in chapter.text.splitlines() if line.strip())
        if document.format == "docx":
            reference = next((b.source_reference for b in blocks if b.source_reference), f"docx:section-{chapter.number}")
        identity = hashlib.sha256(f"{document.format}:{chapter.number}:{reference}".encode()).hexdigest()[:16]
        chapters.append(replace(chapter, id=f"ch_{identity}", source_reference=reference,
            raw_text="\n\n".join(pages[min(hits)-1:max(hits)]) if pages is not None and hits else chapter.text,
            blocks=blocks, word_count=len(chapter.text.split()),
            cleaning_notes=("PDF: limpieza de números y encabezados/pies repetidos solo en bordes de página; el texto original muestra las páginas de origen.",)
                if pages is not None else ()))
        cursor = end
    return replace(document, chapters=tuple(chapters))


def _heading_level(style_name: str) -> int | None:
    match = re.fullmatch(r"(?:heading|t[ií]tulo)\s*([1-9])", style_name, re.I)
    return int(match.group(1)) if match else None


def _clean_pdf_pages(pages: list[str]) -> str:
    """Remove obvious page numbers and repeated first/last lines."""
    return normalize_text("\n\n".join(_pdf_page_texts(pages)))


def _pdf_page_texts(pages: list[str]) -> list[str]:
    page_lines = [
        [line.strip() for line in page.splitlines() if line.strip()]
        for page in pages
    ]
    boundary_lines = [
        line
        for lines in page_lines
        for line in (lines[:1] + lines[-1:])
        if line
    ]
    counts = Counter(boundary_lines)
    repetition_threshold = max(2, (len(page_lines) + 1) // 2)
    repeated = {
        line for line, count in counts.items() if count >= repetition_threshold
    }
    cleaned_pages: list[str] = []
    for lines in page_lines:
        useful_lines = [line for index, line in enumerate(lines)
            if not (index in {0, len(lines)-1} and (line in repeated
                or re.fullmatch(r"(?:p[aá]gina\s+)?\d+", line, re.I)))]
        cleaned_pages.append("\n".join(useful_lines))
    return [normalize_text(re.sub(r"(?<=[a-záéíóúüñ])-\n(?=[a-záéíóúüñ])", "", page,
            flags=re.IGNORECASE)) for page in cleaned_pages]
