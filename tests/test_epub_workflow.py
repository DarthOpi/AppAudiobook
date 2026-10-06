"""V0.9 fixtures and chapter workflows, entirely offline with valid mock WAVs."""

import json
import wave
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest
from ebooklib import epub
from fastapi.testclient import TestClient

from smart_audiobook.application import AudiobookApplicationService
from smart_audiobook.book_manifest import book_statistics, select_range
from smart_audiobook.document_loaders import DocumentLoadError, load_document
from smart_audiobook.epub_importer import html_blocks, validate_epub
from smart_audiobook.models import ChapterStatus
from smart_audiobook.review_service import ReviewService
from smart_audiobook.review_store import ReviewProjectStore, ReviewStateError
from smart_audiobook.speaker_identification import SpeakerIdentificationService
from smart_audiobook.speaker_resolvers import RuleBasedSpeakerResolver
from smart_audiobook.tts_providers import VoiceInfo
from smart_audiobook.web import create_app


def make_epub(path, *, bodies=None, spine=(1, 0), nav=False):
    book = epub.EpubBook()
    book.set_identifier("original-test-fixture")
    book.set_title("La prueba — un libro")
    book.set_language("es")
    book.add_author("Autora de prueba")
    book.add_metadata("DC", "publisher", "Editorial de prueba")
    sections = []
    bodies = bodies or [
        '<h1>Chapter 1</h1><p><em>Una noche…</em></p><p>—Hola —dijo Elena.</p>',
        '<h1>Chapter 2</h1><p>Elena esperó.</p><p>—Adiós —respondió Marcos.</p>',
    ]
    for index, body in enumerate(bodies):
        section = epub.EpubHtml(uid=f"section-{index}", title=f"Chapter {index+1}", file_name=f"section-{index}.xhtml", lang="es")
        section.content = body
        sections.append(section)
        book.add_item(section)
    book.toc = tuple(sections)
    navigation = epub.EpubNav()
    book.add_item(navigation)
    book.add_item(epub.EpubNcx())
    book.spine = ([navigation] if nav else []) + [sections[i] for i in spine]
    epub.write_epub(path, book, options={"raise_exceptions": True})
    return path


class FakeTTS:
    provider_id = "fixture"
    model_id = "pcm-test"
    def __init__(self):
        self.calls = []
    def list_voices(self):
        return [VoiceInfo("a", "A", "es"), VoiceInfo("b", "B", "es")]
    def synthesize(self, text, output_path, voice_id, settings=None):
        self.calls.append((text, voice_id))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with wave.open(str(output_path), "wb") as wav:
            wav.setnchannels(1)
            wav.setsampwidth(2)
            wav.setframerate(8000)
            wav.writeframes(b"\x00\x00" * 80)
        return output_path


@pytest.fixture
def workflow(tmp_path):
    provider = FakeTTS()
    application = AudiobookApplicationService(lambda _use: SpeakerIdentificationService(RuleBasedSpeakerResolver()))
    review = ReviewService(ReviewProjectStore(tmp_path / "work"), provider, application)
    source = make_epub(tmp_path / "book.epub", spine=(0, 1))
    project = review.books.import_book(source)
    return review, provider, project


def test_valid_epub_metadata_spine_and_multiple_chapters(tmp_path):
    document = load_document(make_epub(tmp_path / "test.epub"))
    assert document.title == "La prueba — un libro"
    assert document.author == "Autora de prueba"
    assert document.language == "es"
    assert document.publisher == "Editorial de prueba"
    assert document.identifier == "original-test-fixture"
    assert document.format == "epub" and document.chapter_count == 2
    assert document.chapters[0].title == "Chapter 2"
    assert "Adiós" in document.chapters[0].text
    assert document.chapters[0].source_reference.startswith("section-1.xhtml")
    assert document.full_text == ""  # no giant concatenated book string


def test_epub_without_headings_follows_spine(tmp_path):
    document = load_document(make_epub(tmp_path / "no-headings.epub", bodies=["<p>Uno.</p>", "<p>Dos.</p>"]))
    assert [ch.text for ch in document.chapters] == ["Dos.", "Uno."]


def test_epub_multiple_peer_headings_split_one_item(tmp_path):
    document = load_document(make_epub(tmp_path / "headings.epub", bodies=[
        "<h1>Uno</h1><p>Primero.</p><h1>Dos</h1><p>Segundo.</p>"], spine=(0,)))
    assert [ch.title for ch in document.chapters] == ["Uno", "Dos"]
    assert len({ch.id for ch in document.chapters}) == 2


def test_duplicate_titles_have_stable_different_ids(tmp_path):
    path = make_epub(tmp_path / "duplicate.epub", bodies=["<h1>Uno</h1><p>A.</p>", "<h1>Uno</h1><p>B.</p>"])
    first = load_document(path)
    second = load_document(path)
    assert [ch.id for ch in first.chapters] == [ch.id for ch in second.chapters]
    assert len({ch.id for ch in first.chapters}) == 2


def test_missing_spine_falls_back_to_text_sections(tmp_path):
    document = load_document(make_epub(tmp_path / "no-spine.epub", spine=()))
    assert len(document.chapters) == 2
    assert "Hola" in document.chapters[0].text


def test_no_html_headings_uses_chapter_patterns(tmp_path):
    document = load_document(make_epub(tmp_path / "patterns.epub", bodies=[
        "<p>Chapter 1</p><p>Primero.</p><p>Chapter 2</p><p>Segundo.</p>"], spine=(0,)))
    assert [ch.title for ch in document.chapters] == ["Chapter 1", "Chapter 2"]


def test_malformed_epub_is_a_domain_error(tmp_path):
    source = tmp_path / "bad.epub"
    source.write_bytes(b"not a zip")
    with pytest.raises(DocumentLoadError):
        load_document(source)


@pytest.mark.parametrize("name", ["../../escape.txt", "/absolute.txt", "C:/escape.txt", "x\\escape.txt", "%2e%2e/escape.txt"])
def test_epub_rejects_unsafe_paths_without_extraction(tmp_path, name):
    path = make_epub(tmp_path / "unsafe.epub")
    with ZipFile(path, "a") as archive:
        info = ZipInfo("placeholder")
        info.filename = name  # preserve raw backslashes even on Windows
        archive.writestr(info, "unsafe")
    with pytest.raises(ValueError, match="ruta"):
        validate_epub(path)
    assert not (tmp_path / "escape.txt").exists()


def test_epub_rejects_expansion_and_external_entities(tmp_path):
    path = make_epub(tmp_path / "large.epub")
    with ZipFile(path, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("large.txt", "a" * 1_000_000)
    with pytest.raises(ValueError, match="límites"):
        validate_epub(path)
    path = make_epub(tmp_path / "entities.epub")
    with ZipFile(path, "a") as archive:
        archive.writestr("malicious.xml", '<!DOCTYPE foo [<!ENTITY x SYSTEM "file:///private">]><foo>&x;</foo>')
    with pytest.raises(ValueError, match="entidades"):
        validate_epub(path)


def test_epub_rejects_encryption(tmp_path):
    path = make_epub(tmp_path / "protected.epub")
    with ZipFile(path, "a") as archive:
        archive.writestr("META-INF/encryption.xml", "<encryption/>")
    with pytest.raises(ValueError, match="cifrado"):
        validate_epub(path)


def test_html_sanitization_emphasis_scene_breaks_and_unicode():
    blocks, raw, notes = html_blocks(('<html><head><script>evil()</script></head><body>'
        '<p hidden>Invisible</p><nav>Navigation</nav><p style="display: none">Oculto</p>'
        '<p><em>“Hola…”</em> — <strong>María</strong>&nbsp;sonrió.</p><hr/><p>Fin.</p></body></html>').encode())
    text = "\n".join(b.text for b in blocks)
    assert "evil" not in text and "Invisible" not in text and "Navigation" not in text and "Oculto" not in text
    assert "“Hola…”" in text and "María" in text
    assert blocks[0].emphasis == ("bold", "italic")
    assert any(b.scene_break for b in blocks)
    assert notes and "Invisible" in raw


def test_navigation_is_indexed_but_not_narrated(tmp_path):
    document = load_document(make_epub(tmp_path / "nav.epub", nav=True))
    assert document.chapters[0].section_kind == "table_of_contents"
    assert not document.chapters[0].narrate
    assert document.chapters[1].selected_for_processing


def test_import_does_not_analyze_or_call_gemini_and_loads_lazily(workflow):
    review, _, project = workflow
    assert project.analysis.chapters == ()
    directory = review.store.directory(project.processing_id)
    manifest = json.loads((directory / "project.json").read_text(encoding="utf-8"))
    assert manifest["schema_version"] == 3
    assert manifest["book"]["chapter_count"] == 2
    assert "normalized_text" not in json.dumps(manifest)
    loaded = review.load(project.processing_id)
    chapter = loaded.analysis.document.chapters[0]
    assert chapter.text == "" and chapter.text_path.is_file()
    assert "Hola" in review.books.preview(project.processing_id, 1).text


def test_range_selection_and_invalid_ranges(workflow):
    review, _, project = workflow
    assert select_range("1-2,2", {1, 2}) == {1, 2}
    with pytest.raises(ValueError):
        select_range("1-999999999", {1, 2})
    project = review.books.select_expression(project.processing_id, "2")
    assert [ch.number for ch in project.analysis.document.chapters if ch.selected_for_processing] == [2]
    with pytest.raises(ValueError):
        review.books.select_expression(project.processing_id, "3")


def test_statistics_and_llm_estimate_are_local(workflow):
    review, _, project = workflow
    with patch("smart_audiobook.gemini_provider.GeminiProvider.generate_structured", side_effect=AssertionError("network")):
        stats = review.books.statistics(project.processing_id, detailed=True)
    assert stats["selected_words"] == stats["total_words"] > 0
    assert stats["estimated_dialogues"] == 2
    assert stats["estimated_llm_calls_upper_bound"] == 0
    with pytest.raises(ValueError):
        book_statistics(project.analysis.document, words_per_minute=-1)


def test_partial_analysis_and_later_block_preserve_registry(workflow):
    review, _, project = workflow
    review.books.select_expression(project.processing_id, "1")
    first = review.books.analyze(project.processing_id, use_llm=False)
    assert [ch.number for ch in first.analysis.chapters] == [1]
    assert first.analysis.document.chapters[1].status == ChapterStatus.NOT_ANALYZED
    first_segments = first.analysis.chapters[0].segments
    registry_id = first.analysis.registry.require("Elena").id
    review.edit_alias(project.processing_id, "Elena", "Luz", 2)
    review.books.select_expression(project.processing_id, "2")
    second = review.books.analyze(project.processing_id, use_llm=False)
    assert second.analysis.chapters[0].segments == first_segments
    assert second.analysis.registry.require("Elena").id == registry_id
    assert second.analysis.registry.find("Luz", 1) is None
    assert second.analysis.registry.find("Luz", 2).canonical_name == "Elena"
    assert len(second.analysis.chapters) == 2


def test_reanalysis_of_one_chapter_keeps_manual_overrides(workflow):
    review, _, project = workflow
    project = review.books.analyze(project.processing_id, use_llm=False)
    dialogue = next(s for s in project.segments if s.type == "dialogue" and s.chapter == 1)
    project = review.update_speakers(project.processing_id, {dialogue.id: "Marcos"})
    later = project.analysis.chapters[1]
    project = review.books.analyze(project.processing_id, {1}, reanalyze=True, use_llm=False)
    assert project.analysis.chapters[1] == later
    corrected = next(s for s in project.segments if s.id == dialogue.id)
    assert corrected.speaker == "Marcos" and corrected.resolution_method == "manual"


def test_selected_source_not_reading_unselected_chapter(workflow):
    review, _, project = workflow
    chapter = project.analysis.document.chapters[1]
    chapter.text_path.unlink()
    # Even missing unselected source content should not affect selected analysis.
    project = review.books.analyze(project.processing_id, {1}, use_llm=False)
    assert [ch.number for ch in project.analysis.chapters] == [1]


def test_generate_continue_resume_regenerate_and_clear_one_cache(workflow):
    review, provider, project = workflow
    project = review.books.analyze(project.processing_id, {1}, use_llm=False)
    first = review.generate(project.processing_id, {1})
    first_file = first.output.chapter_files[0]
    first_bytes = first_file.read_bytes()
    calls = len(provider.calls)
    review.generate(project.processing_id, {1})
    assert len(provider.calls) == calls
    project = review.books.analyze(project.processing_id, {2}, use_llm=False)
    second = review.generate(project.processing_id, {2})
    assert len(second.output.chapter_files) == 2
    assert first_file.read_bytes() == first_bytes
    before = len(provider.calls)
    review.generate(project.processing_id, {2}, force_regenerate=True)
    assert len(provider.calls) > before
    assert first_file.read_bytes() == first_bytes
    root = review.store.directory(project.processing_id) / "cache" / "chapters"
    ch1, ch2 = project.analysis.document.chapters
    assert (root / ch1.id).is_dir() and (root / ch2.id).is_dir()
    review.books.clear_chapter_cache(project.processing_id, 2)
    assert (root / ch1.id).is_dir() and not (root / ch2.id).exists()
    before = len(provider.calls)
    review.generate(project.processing_id, {2})
    assert len(provider.calls) > before  # clearing invalidates completed generation too


def test_manifest_source_path_traversal_is_rejected(workflow):
    review, _, project = workflow
    path = review.store.directory(project.processing_id) / "project.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["chapters"][0]["content_reference"] = "../../escape.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ReviewStateError):
        review.load(project.processing_id)


def test_migrate_v08_preserves_manual_segments_voices_and_aliases(workflow):
    review, _, project = workflow
    project = review.books.analyze(project.processing_id, {1}, use_llm=False)
    segment = next(s for s in project.segments if s.type == "dialogue")
    review.update_speakers(project.processing_id, {segment.id: "Elena"})
    project = review.select_voice(project.processing_id, "Elena", "b")
    review.edit_alias(project.processing_id, "Elena", "Luz", 2)
    directory = review.store.directory(project.processing_id)
    path = directory / "analysis.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    for chapter in payload["chapters"]:
        data = json.loads((directory / chapter.pop("analysis_reference")).read_text(encoding="utf-8"))
        chapter["segments"] = data["segments"]
    payload["document"]["full_text"] = "Texto anterior."
    path.write_text(json.dumps(payload), encoding="utf-8")
    (directory / "project.json").unlink()
    migrated = review.load(project.processing_id)
    review.store.save(migrated)
    loaded = review.load(project.processing_id)
    assert (directory / "project.json").is_file()
    assert loaded.dialogues[0].resolution_method == "manual"
    assert loaded.voice_assignments["Elena"] == "b"
    assert loaded.analysis.registry.find("Luz", 1) is None
    assert loaded.analysis.registry.find("Luz", 2).canonical_name == "Elena"


def test_reanalysis_changes_mark_later_chapters_without_destroying_them(workflow):
    from smart_audiobook.models import SpeakerResolution
    class OtherRule:
        def resolve(self, context):
            return SpeakerResolution("Otro", .99, "rule")
    review, _, project = workflow
    project = review.books.analyze(project.processing_id, use_llm=False)
    later = project.analysis.chapters[1]
    project = review.books.analyze(project.processing_id, {1}, reanalyze=True,
        speaker_service=SpeakerIdentificationService(OtherRule()))
    assert project.analysis.chapters[1] == later
    assert project.analysis.document.chapters[1].consistency_warning


def test_imported_source_files_are_not_rewritten_on_review(workflow):
    review, _, project = workflow
    project = review.books.analyze(project.processing_id, use_llm=False)
    source = project.analysis.document.chapters[0].text_path
    before = source.stat().st_mtime_ns
    with patch("smart_audiobook.models.TextSegment.as_dict", side_effect=AssertionError("unchanged analysis was serialized")):
        review.edit_alias(project.processing_id, "Elena", "Luz", 2)
    assert source.stat().st_mtime_ns == before


def test_docx_preserves_emphasis_and_source_reference(tmp_path):
    from docx import Document
    document = Document()
    document.add_heading("Capítulo 1", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Importante.").italic = True
    source = tmp_path / "emphasis.docx"
    document.save(source)
    imported = load_document(source)
    assert any("italic" in b.emphasis for b in imported.chapters[0].blocks)
    assert imported.chapters[0].source_reference.startswith("docx:")


def test_pdf_boundary_cleaning_does_not_remove_repeated_narrative():
    from smart_audiobook.document_loaders import _clean_pdf_pages
    text = _clean_pdf_pages(["HEADER\nPrimero.\nHEADER\n1", "HEADER\nSegundo.\n2"])
    assert text.count("HEADER") == 1
    assert "Primero." in text and "Segundo." in text


def test_scene_break_metadata_and_pause_are_not_synthesized(workflow):
    from smart_audiobook.audio import generate_audiobook
    from smart_audiobook.segmenter import segment_text
    review, provider, project = workflow
    segments = segment_text("Antes.\n***\nDespués.")
    assert len(segments) == 2 and segments[1].scene_break_before
    output = review.store.directory(project.processing_id) / "scene.wav"
    generate_audiobook(segments, output, {"Narrator": "a"}, provider)
    assert [text for text, _ in provider.calls] == ["Antes.", "Después."]
    with wave.open(str(output)) as wav:
        assert wav.getnframes()/wav.getframerate() >= 1


def test_startup_recovers_interrupted_chapters(workflow):
    review, provider, project = workflow
    project.analysis = replace(project.analysis, document=replace(project.analysis.document,
        chapters=tuple(replace(ch, status=ChapterStatus.ANALYZING) for ch in project.analysis.document.chapters)))
    review.store.save(project)
    with TestClient(create_app(review_service=review, tts_provider=provider)) as client:
        status = client.get(f"/books/{project.processing_id}/status").json()
        assert not status["busy"] and status["failed"]
        response = client.post(f"/books/{project.processing_id}/analyze", data={})
        assert response.status_code == 200


def test_manual_edits_do_not_race_running_analysis(workflow):
    review, _, project = workflow
    project = review.books.analyze(project.processing_id, use_llm=False)
    dialogue = project.dialogues[0]
    project.analysis = replace(project.analysis, document=replace(project.analysis.document,
        chapters=tuple(replace(ch, status=ChapterStatus.ANALYZING) for ch in project.analysis.document.chapters)))
    review.store.save(project)
    with pytest.raises(ReviewStateError, match="termine"):
        review.update_speakers(project.processing_id, {dialogue.id: "Manual"})


def test_voice_change_invalidates_only_affected_completed_audio(workflow):
    review, _, project = workflow
    project = review.books.analyze(project.processing_id, use_llm=False)
    review.generate(project.processing_id)
    project = review.select_voice(project.processing_id, "Elena", "a")
    assert project.analysis.document.chapters[0].status == ChapterStatus.READY_FOR_AUDIO
    assert project.analysis.document.chapters[1].status == ChapterStatus.COMPLETED


def test_cli_incremental_commands_use_same_workflow(tmp_path, capsys):
    from smart_audiobook.cli import main
    source = make_epub(tmp_path / "cli.epub", spine=(0, 1))
    root = str(tmp_path / "cli-work")
    with patch("smart_audiobook.cli.build_tts_provider", return_value=FakeTTS()):
        assert main(["import", str(source), "--work-root", root]) == 0
        imported = json.loads(capsys.readouterr().out)
        identity = imported["project_id"]
        assert imported["analyzed_chapters"] == 0
        assert main(["analyze", identity, "--work-root", root, "--chapters", "2", "--no-llm"]) == 0
        analyzed = json.loads(capsys.readouterr().out)
        assert analyzed["analyzed_chapters"] == 1
        assert main(["generate", identity, "--work-root", root, "--chapters", "2"]) == 0
        generated = json.loads(capsys.readouterr().out)
        assert Path(generated["audio"]).is_file()


def test_web_import_selection_preview_partial_analysis_and_search(tmp_path):
    provider = FakeTTS()
    source = make_epub(tmp_path / "web.epub", spine=(0, 1))
    with TestClient(create_app(output_root=tmp_path / "web-work", tts_provider=provider)) as client:
        response = client.post("/upload", files={"document": ("web.epub", source.read_bytes(), "application/epub+zip")})
        assert response.status_code == 200 and "BOOK IMPORTED" in response.text
        identity = response.url.path.split("/")[-1]
        assert "0 / 2" in response.text
        response = client.post(f"/books/{identity}/selection", data={"action": "range", "from": "1", "to": "1"})
        assert response.status_code == 200
        preview = client.get(f"/books/{identity}/chapters/1")
        assert "IMPORT PREVIEW" in preview.text and "Una noche…" in preview.text
        response = client.post(f"/books/{identity}/analyze", data={})
        assert response.status_code == 200
        status = client.get(f"/books/{identity}/status").json()
        assert status["analyzed"] == 1 and status["audio_generated"] == 0
        assert client.get(f"/books/{identity}?q=2").status_code == 200
        assert client.get(f"/review/{identity}").status_code == 200
        generated = client.post(f"/review/{identity}/generate")
        assert generated.status_code == 200 and "AUDIOBOOK READY" in generated.text
        status = client.get(f"/books/{identity}/status").json()
        assert status["analyzed"] == 1 and status["audio_generated"] == 1
        assert client.post(f"/books/{identity}/selection", data={"action": "range", "from": "2", "to": "2"}).status_code == 200
        assert client.post(f"/books/{identity}/analyze", data={}).status_code == 200
        generated = client.post(f"/review/{identity}/generate")
        assert generated.status_code == 200
        status = client.get(f"/books/{identity}/status").json()
        assert status["analyzed"] == 2 and status["audio_generated"] == 2
