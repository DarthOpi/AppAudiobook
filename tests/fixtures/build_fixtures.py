"""Rebuild the committed PDF and DOCX integration fixtures."""

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

FIXTURE_DIRECTORY = Path(__file__).resolve().parent


def build_docx() -> None:
    document = Document()
    document.core_properties.title = "El faro perdido"
    document.add_heading("El faro perdido", level=0)
    document.add_heading("Prólogo", level=1)
    document.add_paragraph("La niebla cubría el puerto.")
    document.add_heading("Capítulo 1: La señal", level=1)
    document.add_paragraph("Lucía encendió la lámpara.")
    document.add_paragraph("—¿La ves? —preguntó Lucía.")
    document.add_heading("Capítulo 2: El regreso", level=1)
    document.add_paragraph("Mateo llegó antes del amanecer.")
    document.add_paragraph("—He vuelto —dijo Mateo.")
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(11)
    document.save(FIXTURE_DIRECTORY / "sample_book.docx")


def build_pdf() -> None:
    output_path = FIXTURE_DIRECTORY / "sample_book.pdf"
    styles = getSampleStyleSheet()
    story = [
        Paragraph("El reloj de arena", styles["Title"]),
        Spacer(1, 10 * mm),
        Paragraph("Capítulo 1: La carta", styles["Heading1"]),
        Paragraph("Elena encontró una carta sobre la mesa.", styles["BodyText"]),
        Paragraph("—No puede ser —dijo Elena.", styles["BodyText"]),
        PageBreak(),
        Paragraph("Capítulo 2: El viaje", styles["Heading1"]),
        Paragraph("Tomás preparó la mochila al amanecer.", styles["BodyText"]),
        Paragraph("—Nos vamos —respondió Tomás.", styles["BodyText"]),
    ]
    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        title="El reloj de arena",
        author="Smart Audiobook test suite",
    )
    pdf.build(story)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("format", choices=("docx", "pdf", "all"), default="all", nargs="?")
    selected = parser.parse_args().format
    if selected in {"docx", "all"}:
        build_docx()
    if selected in {"pdf", "all"}:
        build_pdf()
